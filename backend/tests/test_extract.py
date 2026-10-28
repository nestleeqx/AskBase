import io

import pytest
from docx import Document as DocxDocument

from app.services.chunking import MAX_CHARS, embedding_input, make_chunks, split_long
from app.services.extract import (
    ExtractError,
    extract_docx,
    extract_file,
    extract_html,
    extract_txt,
    kind_of,
    looks_like_heading,
    table_units,
)
from app.services.seed import SEED_BY_KEY


def text_of(units):
    return "\n".join(t for _, t in units)


def test_kind_by_extension():
    assert kind_of("Прайс.PDF") == "pdf"
    assert kind_of("врачи.docx") == "docx"
    assert kind_of("faq.txt") == "txt"
    with pytest.raises(ExtractError, match="DOCX"):
        kind_of("старый.doc")
    with pytest.raises(ExtractError, match="PDF, DOCX и TXT"):
        kind_of("таблица.xlsx")


@pytest.mark.parametrize(
    ("line", "expected"),
    [
        ("Хирургия", True),
        ("Исправление прикуса (ортодонтия)", True),
        ("Как подготовиться к удалению зуба?", False),
        ("4.01 Удаление зуба простое 3 500", False),
        ("Воскресенье — выходной.", False),
        ("нашей клинике)", False),
    ],
)
def test_heading_heuristic(line, expected):
    assert looks_like_heading(line) is expected


def test_seed_pdf_rows_and_headings():
    seed = SEED_BY_KEY["prajs"]
    doc = extract_file("pdf", seed.path.read_bytes(), seed.file, max_pages=20)
    headings = [t for k, t in doc.units if k == "h"]
    assert "Хирургия" in headings and "Имплантация" in headings
    body = text_of(doc.units)
    # Шапка таблицы повторялась в каждом разделе и колонтитул на каждой странице — их нет
    assert "Код Услуга" not in body
    assert "стр." not in body
    # Цена, перенесённая на отдельную строку, приклеена к своей услуге
    assert any(t.startswith("2.01 ") and t.endswith("5 900") for _, t in doc.units)
    assert any(t.startswith("5.01 ") and t.endswith("54 900") for _, t in doc.units)
    assert doc.pages == 3


def test_seed_docx_table_rows_keep_column_names():
    seed = SEED_BY_KEY["vrachi"]
    doc = extract_file("docx", seed.path.read_bytes(), seed.file, max_pages=20)
    row = next(t for _, t in doc.units if "Карпова" in t)
    assert row.startswith("Врач: Карпова Елена Дмитриевна; Специализация: Ортодонт")
    assert "Дни и часы приёма: Пн, Ср 12:00–20:00; Пт 9:00–14:00" in row
    assert doc.title == "Врачи и график работы клиники «Белая линия»"


def test_docx_headings_and_lists():
    d = DocxDocument()
    d.add_heading("Оплата", level=1)
    d.add_paragraph("Принимаем карты.")
    d.add_paragraph("наличные", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    doc = extract_docx(buf.getvalue(), "файл")
    assert doc.units == [("h", "Оплата"), ("p", "Принимаем карты."), ("p", "• наличные")]


def test_txt_caps_headings_and_qa_paragraphs():
    doc = extract_txt(
        "ОПЛАТА\n\nМожно ли в рассрочку?\nДа, до 12 месяцев.\n\nГарантии\n\nПломбы — 1 год.",
        "faq",
    )
    assert doc.units == [
        ("h", "ОПЛАТА"),
        ("p", "Можно ли в рассрочку?\nДа, до 12 месяцев."),
        ("h", "Гарантии"),
        ("p", "Пломбы — 1 год."),
    ]


def test_html_skips_menu_and_footer():
    html = """<html><head><title>Клиника</title></head><body>
      <header><nav>Услуги Врачи Контакты</nav></header>
      <main><h1>Правила</h1><p>Отмена за <b>24 часа</b>.</p>
        <ul><li><p>паспорт</p></li><li>снимки</li></ul>
        <table><tr><th>Услуга</th><th>Цена</th></tr><tr><td>Чистка</td><td>5 900</td></tr></table>
      </main><footer>© Клиника</footer></body></html>"""
    doc = extract_html(html, "страница")
    assert doc.title == "Правила"
    assert doc.units == [
        ("h", "Правила"),
        ("p", "Отмена за 24 часа."),
        ("p", "• паспорт"),
        ("p", "• снимки"),
        ("p", "Услуга: Чистка; Цена: 5 900"),
    ]


def test_table_without_header_row():
    assert table_units([["Только", "одна строка"]]) == [("p", "Только | одна строка")]


def test_broken_files_give_user_messages():
    with pytest.raises(ExtractError, match="PDF"):
        extract_file("pdf", b"not a pdf", "x.pdf", max_pages=20)
    with pytest.raises(ExtractError, match="DOCX"):
        extract_file("docx", b"not a zip", "x.docx", max_pages=20)


def test_pdf_page_limit():
    seed = SEED_BY_KEY["prajs"]
    with pytest.raises(ExtractError, match="больше 2 страниц"):
        extract_file("pdf", seed.path.read_bytes(), seed.file, max_pages=2)


def test_cp1251_txt():
    doc = extract_file("txt", "Цены и скидки".encode("cp1251"), "a.txt", max_pages=20)
    assert doc.units == [("h", "Цены и скидки")]


# --- Нарезка ---


def test_split_long_by_sentences():
    text = " ".join(f"Предложение номер {i} про лечение зубов." for i in range(60))
    parts = split_long(text)
    assert len(parts) > 1
    assert all(len(p) <= MAX_CHARS for p in parts)
    assert " ".join(parts) == text


def test_chunks_respect_limit_and_sections():
    units = [("h", "Раздел А")]
    units += [("p", f"Абзац {i}: " + "текст " * 30) for i in range(12)]
    units += [("h", "Раздел Б"), ("p", "Короткий абзац раздела Б.")]
    drafts = make_chunks(units)
    assert all(len(d.text) <= MAX_CHARS for d in drafts)
    assert drafts[0].heading == "Раздел А"
    # Продолжение раздела помнит свой заголовок, хотя строки заголовка в нём нет
    assert drafts[1].heading == "Раздел А" and not drafts[1].text.startswith("Раздел А")
    assert drafts[-1].text.endswith("Короткий абзац раздела Б.")


def test_small_sections_merge():
    units = [("h", "Дети"), ("p", "С 2 лет."), ("h", "Гарантии"), ("p", "Пломбы — 1 год.")]
    drafts = make_chunks(units)
    assert len(drafts) == 1
    assert drafts[0].text == "Дети\nС 2 лет.\nГарантии\nПломбы — 1 год."


def test_seed_documents_chunk_sizes():
    for seed in SEED_BY_KEY.values():
        kind = "html" if seed.kind == "url" else seed.kind
        doc = extract_file(kind, seed.path.read_bytes(), seed.file, max_pages=20)
        drafts = make_chunks(doc.units)
        assert drafts, seed.key
        assert all(len(d.text) <= MAX_CHARS for d in drafts), seed.key


def test_embedding_input_adds_context():
    drafts = make_chunks([("h", "Хирургия")] + [("p", "строка " * 40)] * 5)
    tail = drafts[-1]
    assert embedding_input("Прайс-лист", tail).startswith("Прайс-лист\nХирургия\n")
