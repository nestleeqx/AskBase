import io

import pytest
from docx import Document as DocxDocument
from sqlalchemy import func, select

from app.db import SessionLocal
from app.embeddings.mock import MockEmbeddings
from app.models import Chunk, Document
from app.services import fetch as fetch_module
from app.services.seed import ensure_seed
from tests.conftest import OTHER

FAQ_TXT = (
    "ПАРКОВКА\n\nЕсть ли у вас парковка?\nДа, бесплатная парковка во дворе на 6 машин.\n\n"
    "ВОДА\n\nЕсть ли кулер?\nДа, в холле стоит кулер с водой."
).encode()


async def upload(client, name="parking.txt", data=FAQ_TXT):
    return await client.post("/api/documents", files={"file": (name, data)})


async def test_seed_documents_are_ready(client):
    resp = await client.get("/api/documents")
    assert resp.status_code == 200
    body = resp.json()
    docs = body["documents"]
    assert [d["title"] for d in docs] == [
        "Прайс-лист",
        "Врачи и график работы",
        "Частые вопросы пациентов",
        "Правила записи на приём",
    ]
    assert all(d["status"] == "ready" and d["chunks_count"] > 0 for d in docs)
    assert not any(d["own"] for d in docs)
    pravila = docs[3]
    assert pravila["kind"] == "url"
    assert pravila["source_url"] == "https://clinic.example/clinic/pravila-zapisi.html"
    assert body["limits"]["guest_docs_left"] == 3


async def test_seed_is_not_reindexed_on_restart(client, redis):
    async with SessionLocal() as s:
        before = list(await s.scalars(select(Chunk.id).order_by(Chunk.id)))
    async with SessionLocal() as s:
        from app.main import app

        await ensure_seed(SessionLocal, app.state.embedder, redis)
    async with SessionLocal() as s:
        assert list(await s.scalars(select(Chunk.id).order_by(Chunk.id))) == before


async def test_new_embedding_model_reembeds_documents(client, redis):
    class OtherMock(MockEmbeddings):
        model = "other-hash"

    other = OtherMock()
    await ensure_seed(SessionLocal, other, redis)
    async with SessionLocal() as s:
        models = set(await s.scalars(select(Chunk.embedding_model).distinct()))
    assert models == {"mock:other-hash"}
    # Возвращаем как было, чтобы остальные тесты искали обычной моделью
    from app.main import app

    await ensure_seed(SessionLocal, app.state.embedder, redis)


async def test_original_file_of_seed_document(client):
    docs = (await client.get("/api/documents")).json()["documents"]
    resp = await client.get(f"/api/documents/{docs[0]['id']}/file")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content.startswith(b"%PDF")
    # У страницы сайта оригинала-файла нет — у неё есть ссылка
    assert (await client.get(f"/api/documents/{docs[3]['id']}/file")).status_code == 404


async def test_upload_txt_indexes_and_searches(client):
    resp = await upload(client)
    assert resp.status_code == 201
    doc = resp.json()
    assert doc["status"] == "ready" and doc["own"] and doc["chunks_count"] == 1
    assert doc["index_ms"] is not None and doc["index_ms"] < 60_000
    assert doc["title"] == "parking"

    chunks = (await client.get(f"/api/documents/{doc['id']}/chunks")).json()["chunks"]
    assert "парковка во дворе" in chunks[0]["text"]

    found = (
        await client.post("/api/search", json={"question": "Где оставить машину, есть парковка?"})
    ).json()
    assert found["results"][0]["document"]["id"] == doc["id"]
    assert found["results"][0]["keyword_rank"] == 1


async def test_own_documents_are_private(client):
    doc = (await upload(client)).json()
    other = {"X-Visitor-Id": OTHER}
    titles = [
        d["title"] for d in (await client.get("/api/documents", headers=other)).json()["documents"]
    ]
    assert "parking" not in titles
    assert (
        await client.get(f"/api/documents/{doc['id']}/chunks", headers=other)
    ).status_code == 404
    found = (await client.post("/api/search", json={"question": "парковка"}, headers=other)).json()
    assert all(r["document"]["id"] != doc["id"] for r in found["results"])
    assert (await client.delete(f"/api/documents/{doc['id']}", headers=other)).status_code == 404


async def test_guest_limit_and_delete(client):
    ids = [(await upload(client, f"doc{i}.txt")).json()["id"] for i in range(3)]
    resp = await upload(client, "doc4.txt")
    assert resp.status_code == 409
    assert "до 3 своих документов" in resp.json()["detail"]

    assert (await client.delete(f"/api/documents/{ids[0]}")).status_code == 204
    assert (await upload(client, "doc4.txt")).status_code == 201
    async with SessionLocal() as s:
        assert (
            await s.scalar(
                select(func.count()).select_from(Chunk).where(Chunk.document_id == ids[0])
            )
            == 0
        )


async def test_clinic_documents_cannot_be_deleted(client):
    doc = (await client.get("/api/documents")).json()["documents"][0]
    resp = await client.delete(f"/api/documents/{doc['id']}")
    assert resp.status_code == 403
    assert "общая база" in resp.json()["detail"]


async def test_docx_upload(client):
    d = DocxDocument()
    d.add_heading("Акции", level=1)
    d.add_paragraph("В ноябре скидка 15% на профессиональную гигиену.")
    buf = io.BytesIO()
    d.save(buf)
    doc = (await upload(client, "акции.docx", buf.getvalue())).json()
    assert doc["status"] == "ready" and doc["title"] == "Акции" and doc["kind"] == "docx"


@pytest.mark.parametrize(
    ("name", "data", "status", "message"),
    [
        ("old.doc", b"x", 415, "DOCX"),
        ("table.xlsx", b"x", 415, "PDF, DOCX и TXT"),
        ("big.txt", b"a" * (2 * 1024 * 1024 + 1), 413, "больше 2 МБ"),
        ("empty.txt", b"", 422, "пустой"),
    ],
)
async def test_rejected_uploads(client, name, data, status, message):
    resp = await upload(client, name, data)
    assert resp.status_code == status
    assert message in resp.json()["detail"]


async def test_unreadable_file_keeps_error_in_status(client):
    resp = await upload(client, "scan.pdf", b"%PDF-1.4 broken")
    assert resp.status_code == 201
    doc = resp.json()
    assert doc["status"] == "error"
    assert "PDF" in doc["error"]


async def test_url_upload(client, monkeypatch):
    async def fake_fetch(url, *, allow_private, max_bytes):
        assert url == "https://example.com/akcii"
        html = (
            "<html><body><main><h1>Акции клиники</h1>"
            "<p>Скидка 15% на гигиену в ноябре.</p></main></body></html>"
        )
        return html.encode(), "html"

    monkeypatch.setattr("app.api.documents.fetch_page", fake_fetch)
    resp = await client.post("/api/documents/url", json={"url": "example.com/akcii"})
    doc = resp.json()
    assert doc["status"] == "ready" and doc["kind"] == "url" and doc["title"] == "Акции клиники"
    assert doc["source_url"] == "https://example.com/akcii"


async def test_private_urls_are_blocked():
    with pytest.raises(fetch_module.ExtractError, match="внутренней сети"):
        await fetch_module.fetch_page("http://127.0.0.1:8000/", allow_private=False, max_bytes=1000)
    with pytest.raises(fetch_module.ExtractError, match="http"):
        await fetch_module.fetch_page("ftp://example.com/", allow_private=False, max_bytes=1000)


async def test_search_finds_clinic_facts(client):
    resp = await client.post("/api/search", json={"question": "Кто ставит имплантаты, Мельников?"})
    assert resp.status_code == 200
    body = resp.json()
    assert len(body["results"]) == 6
    assert body["embedding_model"] == "mock:hashed-words"
    top = body["results"][0]
    assert top["rank"] == 1 and "Мельников" in top["chunk"]["text"]
    assert top["document"]["title"] == "Врачи и график работы"
    assert 0 < top["similarity"] <= 1


async def test_search_validation(client):
    resp = await client.post("/api/search", json={"question": "а" * 501})
    assert resp.status_code == 422
    assert "500 символов" in resp.json()["detail"]
    assert (await client.post("/api/search", json={"question": "   "})).status_code == 422


async def test_search_rate_limit(client, monkeypatch):
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "searches_per_hour", 2)
    for _ in range(2):
        assert (await client.post("/api/search", json={"question": "цены"})).status_code == 200
    resp = await client.post("/api/search", json={"question": "цены"})
    assert resp.status_code == 429
    assert "Лимит демо" in resp.json()["detail"]


async def test_visitor_header_required(client):
    resp = await client.get("/api/documents", headers={"X-Visitor-Id": "x"})
    assert resp.status_code == 400


async def test_cleanup_removes_old_guest_docs(client, monkeypatch):
    from datetime import timedelta

    from app.config import get_settings
    from app.models import now

    doc = (await upload(client)).json()
    async with SessionLocal() as s:
        row = await s.get(Document, doc["id"])
        row.created_at = now() - timedelta(hours=25)
        await s.commit()
    monkeypatch.setattr(get_settings(), "cron_secret", "s3cret")
    assert (await client.get("/api/cron/cleanup")).status_code == 401
    resp = await client.get("/api/cron/cleanup", headers={"Authorization": "Bearer s3cret"})
    assert resp.json() == {"removed": 1}
