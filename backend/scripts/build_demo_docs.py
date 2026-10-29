"""Собирает демо-документы клиники «Белая линия» из Markdown в app/seed/source/.

prajs-list.md      → docs/prajs-list.pdf       (PDF через Chromium)
vrachi-i-grafik.md → docs/vrachi-i-grafik.docx (DOCX через python-docx)
pravila-zapisi.md  → docs/pravila-zapisi.html  (страница сайта клиники, копия — во frontend/public)

faq.txt пишется руками и лежит в docs/ сразу. Скрипт нужен, только если правите тексты:

    uv run --no-project --with playwright==1.64.0 --with python-docx \
        python scripts/build_demo_docs.py
"""

import re
from html import escape
from pathlib import Path

SEED = Path(__file__).resolve().parents[1] / "app" / "seed"
SOURCE, DOCS = SEED / "source", SEED / "docs"
# Сайт клиники отдаёт ту же страницу: на неё ведёт ссылка «Страница» у документа в админке
SITE_PAGE = Path(__file__).resolve().parents[2] / "frontend/public/clinic/pravila-zapisi.html"

CLINIC = "Белая линия"
PHONE = "+7 (495) 200-18-18"
PHONE_TEL = "+74952001818"
CONTACTS = f"Москва, ул. Светлая, 18 · {PHONE}"
# Знак клиники в шапке сайта — контур зуба
TOOTH = (
    "M12.5 9.5c2 0 3.2 1 5.5 1s3.5-1 5.5-1c2.8 0 4.5 2.3 4.5 5.4 0 3.4-1.3 5.6-2.1 8.4-.6 2.1-1 "
    "4.2-2.6 4.2-2.2 0-1.8-6-5.3-6s-3.1 6-5.3 6c-1.6 0-2-2.1-2.6-4.2C9.3 20.5 8 18.3 8 14.9c0-3.1 "
    "1.7-5.4 4.5-5.4z"
)


# --- Разбор Markdown: заголовки, абзацы, списки, таблицы — больше в документах ничего нет ---


def parse(md: str) -> list[tuple[str, object]]:
    blocks: list[tuple[str, object]] = []
    for line in md.splitlines():
        line = line.rstrip()
        if not line:
            continue
        if line.startswith("## "):
            blocks.append(("h2", line[3:]))
        elif line.startswith("# "):
            blocks.append(("h1", line[2:]))
        elif line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r"-+", c) for c in cells):
                continue
            if blocks and blocks[-1][0] == "table":
                blocks[-1][1].append(cells)
            else:
                blocks.append(("table", [cells]))
        elif line.startswith("- "):
            if blocks and blocks[-1][0] == "ul":
                blocks[-1][1].append(line[2:])
            else:
                blocks.append(("ul", [line[2:]]))
        else:
            blocks.append(("p", line))
    return blocks


def to_html(blocks: list[tuple[str, object]]) -> str:
    out = []
    for kind, value in blocks:
        if kind == "table":
            head, *rows = value
            th = "".join(f"<th>{escape(c)}</th>" for c in head)
            body = "".join(
                "<tr>" + "".join(f"<td>{escape(c)}</td>" for c in row) + "</tr>" for row in rows
            )
            out.append(f"<table><thead><tr>{th}</tr></thead><tbody>{body}</tbody></table>")
        elif kind == "ul":
            out.append("<ul>" + "".join(f"<li>{escape(i)}</li>" for i in value) + "</ul>")
        else:
            out.append(f"<{kind}>{escape(value)}</{kind}>")
    return "\n".join(out)


# --- PDF: прайс-лист ---

PRICE_CSS = """
@page { size: A4; margin: 16mm 14mm 18mm; }
* { box-sizing: border-box; }
body { font-family: "Segoe UI", Arial, sans-serif; color: #10302e; font-size: 10.5pt; margin: 0; }
.brand { display: flex; justify-content: space-between; align-items: baseline;
  border-bottom: 2px solid #14b8a6; padding-bottom: 6px; margin-bottom: 14px; }
.brand b { font-size: 15pt; color: #0f766e; letter-spacing: .02em; }
.brand span { color: #4b6463; font-size: 9pt; }
h1 { font-size: 17pt; margin: 0 0 4px; }
h1 + p { color: #4b6463; margin: 0 0 12px; }
h2 { font-size: 12pt; color: #0f766e; margin: 16px 0 6px; break-after: avoid; }
table { width: 100%; border-collapse: collapse; break-inside: auto; }
tr { break-inside: avoid; }
th { text-align: left; font-size: 8.5pt; color: #4b6463; font-weight: 600;
  border-bottom: 1px solid #99d5cf; padding: 4px 6px; }
td { border-bottom: 1px solid #e3efed; padding: 5px 6px; vertical-align: top; }
td:first-child { width: 42px; color: #4b6463; }
td:last-child, th:last-child { width: 92px; text-align: right; white-space: nowrap; }
td:last-child { font-weight: 600; }
p { margin: 8px 0; line-height: 1.45; }
"""


def build_pdf(blocks: list[tuple[str, object]], path: Path) -> None:
    from playwright.sync_api import sync_playwright

    html = (
        f"<!doctype html><html lang=ru><meta charset=utf-8><style>{PRICE_CSS}</style>"
        f'<div class="brand"><b>{CLINIC}</b><span>{escape(CONTACTS)}</span></div>'
        f"{to_html(blocks)}</html>"
    )
    footer = (
        '<div style="font-size:8px;color:#6b8483;width:100%;text-align:center;'
        'font-family:Segoe UI,Arial">Прайс-лист клиники «Белая линия» · стр. '
        '<span class="pageNumber"></span> из <span class="totalPages"></span></div>'
    )
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.set_content(html, wait_until="load")
        page.pdf(
            path=str(path),
            format="A4",
            print_background=True,
            display_header_footer=True,
            header_template="<span></span>",
            footer_template=footer,
            margin={"top": "16mm", "bottom": "18mm", "left": "14mm", "right": "14mm"},
        )
        browser.close()


# --- DOCX: врачи и график ---


def build_docx(blocks: list[tuple[str, object]], path: Path) -> None:
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    teal = RGBColor(0x0F, 0x76, 0x6E)

    for kind, value in blocks:
        if kind in ("h1", "h2"):
            heading = doc.add_heading(value, level=1 if kind == "h1" else 2)
            for run in heading.runs:
                run.font.color.rgb = teal
        elif kind == "table":
            head, *rows = value
            table = doc.add_table(rows=1, cols=len(head))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, head, strict=True):
                cell.text = ""
                cell.paragraphs[0].add_run(text).bold = True
            for row in rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row, strict=True):
                    cell.text = text
            widths = [Cm(4.2), Cm(6.6), Cm(1.6), Cm(4.6)]
            for row in table.rows:
                for cell, width in zip(row.cells, widths, strict=False):
                    cell.width = width
            doc.add_paragraph()
        elif kind == "ul":
            for item in value:
                doc.add_paragraph(item, style="List Bullet")
        else:
            doc.add_paragraph(value)
    doc.core_properties.title = blocks[0][1]
    doc.core_properties.author = f"Клиника «{CLINIC}»"
    doc.save(str(path))


# --- HTML: страница сайта клиники ---

PAGE_CSS = """
:root { --accent: #0f766e; --accent-soft: #e6f6f3; --ink: #10302e; --muted: #4b6463;
  --line: #d7ebe8; --bg: #f7fbfa; }
* { box-sizing: border-box; }
body { margin: 0; font-family: "Segoe UI", Arial, sans-serif; color: var(--ink);
  background: var(--bg); line-height: 1.6; }
header { position: sticky; top: 0; z-index: 10; background: rgba(255, 255, 255, .92);
  backdrop-filter: blur(8px); border-bottom: 1px solid var(--line); }
.bar { display: flex; align-items: center; gap: 24px; height: 64px; }
.logo { display: flex; align-items: center; gap: 10px; color: var(--ink); text-decoration: none;
  line-height: 1.15; }
.logo svg { flex: none; }
.logo b { display: block; font-size: 17px; }
.logo small { display: block; color: var(--muted); font-size: 12px; }
nav { display: flex; gap: 4px; font-size: 14px; }
nav a { padding: 6px 12px; border-radius: 999px; color: var(--muted); text-decoration: none;
  white-space: nowrap; }
nav a:hover { color: var(--ink); background: var(--bg); }
nav a[aria-current] { color: var(--accent); background: var(--accent-soft); font-weight: 600; }
.call { margin-left: auto; display: flex; align-items: center; gap: 14px; }
.phone { color: var(--ink); font-weight: 600; font-size: 14px; text-decoration: none;
  white-space: nowrap; }
.btn { display: inline-flex; align-items: center; gap: 6px; height: 38px; padding: 0 16px;
  border-radius: 10px; background: var(--accent); color: #fff; font-size: 14px; font-weight: 600;
  text-decoration: none; white-space: nowrap; }
.btn:hover { background: #115e59; }
.wrap { max-width: 960px; margin: 0 auto; padding: 0 16px; }
main .wrap { max-width: 760px; }
/* Телефон: логотип и кнопка в одну строку, разделы — под ними, переносятся во вторую строку.
   Медиазапросы — после общих правил .wrap, иначе те перебили бы отступы */
@media (max-width: 760px) {
  .wrap { padding: 0 20px; }
  .bar { flex-wrap: wrap; height: auto; gap: 0; padding-top: 18px; padding-bottom: 16px; }
  .bar nav { order: 3; width: 100%; flex-wrap: wrap; gap: 8px; margin-top: 16px; }
  nav a { padding: 5px 12px; font-size: 13px; background: var(--bg); }
  .phone { display: none; }
  .btn { height: 36px; padding: 0 12px; }
}
@media (max-width: 380px) {
  .logo small { display: none; }
}
main { padding: 32px 0 48px; }
h1 { font-size: 30px; line-height: 1.2; margin: 0 0 20px; }
h2 { font-size: 19px; margin: 28px 0 8px; color: var(--accent); }
p, li { font-size: 16px; }
ul { padding-left: 22px; }
footer { border-top: 1px solid var(--line); color: var(--muted); font-size: 13px; padding: 18px 0; }
"""


def build_page(blocks: list[tuple[str, object]], path: Path) -> None:
    title = blocks[0][1]
    html = f"""<!doctype html>
<html lang="ru">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{escape(title)} — стоматология «{CLINIC}»</title>
<style>{PAGE_CSS}</style>
</head>
<body>
<header><div class="wrap bar">
  <a class="logo" href="#">
    <svg width="36" height="36" viewBox="0 0 36 36" aria-hidden="true">
      <rect width="36" height="36" rx="10" fill="#0f766e"/>
      <path d="{TOOTH}" fill="#fff"/>
    </svg>
    <span><b>{CLINIC}</b><small>стоматологическая клиника</small></span>
  </a>
  <nav aria-label="Разделы сайта">
    <a href="#">Услуги и цены</a><a href="#">Врачи</a>
    <a href="#" aria-current="page">Правила записи</a><a href="#">Контакты</a>
  </nav>
  <div class="call">
    <a class="phone" href="tel:{PHONE_TEL}">{PHONE}</a>
    <a class="btn" href="tel:{PHONE_TEL}">Записаться</a>
  </div>
</div></header>
<main><div class="wrap">
{to_html(blocks)}
</div></main>
<footer><div class="wrap">Стоматология «{CLINIC}» · {escape(CONTACTS)}<br>
Демо-сайт для проекта AskBase: клиника, врачи и цены вымышлены.</div></footer>
</body>
</html>
"""
    path.write_text(html, encoding="utf-8")
    SITE_PAGE.parent.mkdir(parents=True, exist_ok=True)
    SITE_PAGE.write_text(html, encoding="utf-8")


def main() -> None:
    DOCS.mkdir(exist_ok=True)
    build_pdf(parse((SOURCE / "prajs-list.md").read_text("utf-8")), DOCS / "prajs-list.pdf")
    build_docx(
        parse((SOURCE / "vrachi-i-grafik.md").read_text("utf-8")), DOCS / "vrachi-i-grafik.docx"
    )
    build_page(
        parse((SOURCE / "pravila-zapisi.md").read_text("utf-8")), DOCS / "pravila-zapisi.html"
    )
    for f in sorted(DOCS.iterdir()):
        print(f"{f.name:28} {f.stat().st_size / 1024:6.1f} КБ")


if __name__ == "__main__":
    main()
